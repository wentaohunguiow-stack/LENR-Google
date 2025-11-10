"""Excel file handling utilities for RAG system."""

import json
from pathlib import Path
from typing import Dict, List, Optional, Union
import pandas as pd
from openpyxl import load_workbook

from src.utils.logger import setup_logger

logger = setup_logger(__name__)


class ExcelHandler:
    """
    Handler for processing Excel files for RAG ingestion.

    Provides multiple strategies for converting Excel data into
    searchable text format that works well with the RAG system.
    """

    @staticmethod
    def excel_to_markdown(
        file_path: Union[str, Path],
        include_sheet_names: Optional[List[str]] = None
    ) -> str:
        """
        Convert Excel file to Markdown format.

        Args:
            file_path: Path to Excel file (.xlsx or .xls)
            include_sheet_names: List of sheet names to include (None = all sheets)

        Returns:
            Markdown-formatted string representation of the Excel data

        Example:
            >>> handler = ExcelHandler()
            >>> markdown = handler.excel_to_markdown("data.xlsx")
            >>> print(markdown)
        """
        file_path = Path(file_path)
        logger.info(f"Converting {file_path.name} to Markdown")

        # Read all sheets
        excel_file = pd.ExcelFile(file_path)
        sheet_names = include_sheet_names or excel_file.sheet_names

        markdown_parts = [f"# Excel File: {file_path.name}\n"]

        for sheet_name in sheet_names:
            if sheet_name not in excel_file.sheet_names:
                logger.warning(f"Sheet '{sheet_name}' not found, skipping")
                continue

            df = pd.read_excel(excel_file, sheet_name=sheet_name)

            # Add sheet header
            markdown_parts.append(f"\n## Sheet: {sheet_name}\n")

            # Add sheet statistics
            markdown_parts.append(f"**Rows:** {len(df)}, **Columns:** {len(df.columns)}\n")

            # Add the data table
            markdown_parts.append(df.to_markdown(index=False))
            markdown_parts.append("\n")

        return "\n".join(markdown_parts)

    @staticmethod
    def excel_to_structured_text(
        file_path: Union[str, Path],
        include_sheet_names: Optional[List[str]] = None,
        max_rows_per_sheet: Optional[int] = None
    ) -> str:
        """
        Convert Excel to human-readable structured text.

        Better for RAG systems as it provides more context than raw tables.

        Args:
            file_path: Path to Excel file
            include_sheet_names: Sheets to include (None = all)
            max_rows_per_sheet: Maximum rows to process per sheet

        Returns:
            Structured text representation

        Example:
            >>> text = ExcelHandler.excel_to_structured_text("sales.xlsx")
        """
        file_path = Path(file_path)
        logger.info(f"Converting {file_path.name} to structured text")

        excel_file = pd.ExcelFile(file_path)
        sheet_names = include_sheet_names or excel_file.sheet_names

        text_parts = [f"EXCEL FILE: {file_path.name}\n{'=' * 80}\n"]

        for sheet_name in sheet_names:
            if sheet_name not in excel_file.sheet_names:
                continue

            df = pd.read_excel(excel_file, sheet_name=sheet_name)

            if max_rows_per_sheet:
                df = df.head(max_rows_per_sheet)

            text_parts.append(f"\nSHEET: {sheet_name}")
            text_parts.append("-" * 80)
            text_parts.append(f"Total Records: {len(df)}")
            text_parts.append(f"Columns: {', '.join(df.columns)}\n")

            # Convert each row to descriptive text
            for idx, row in df.iterrows():
                text_parts.append(f"\nRecord {idx + 1}:")
                for col in df.columns:
                    value = row[col]
                    if pd.notna(value):  # Skip NaN values
                        text_parts.append(f"  - {col}: {value}")

            text_parts.append("\n" + "=" * 80)

        return "\n".join(text_parts)

    @staticmethod
    def excel_to_json(
        file_path: Union[str, Path],
        include_sheet_names: Optional[List[str]] = None
    ) -> str:
        """
        Convert Excel to JSON format.

        Args:
            file_path: Path to Excel file
            include_sheet_names: Sheets to include (None = all)

        Returns:
            JSON string representation

        Example:
            >>> json_str = ExcelHandler.excel_to_json("data.xlsx")
        """
        file_path = Path(file_path)
        logger.info(f"Converting {file_path.name} to JSON")

        excel_file = pd.ExcelFile(file_path)
        sheet_names = include_sheet_names or excel_file.sheet_names

        result = {
            "filename": file_path.name,
            "sheets": {}
        }

        for sheet_name in sheet_names:
            if sheet_name not in excel_file.sheet_names:
                continue

            df = pd.read_excel(excel_file, sheet_name=sheet_name)
            result["sheets"][sheet_name] = df.to_dict(orient="records")

        return json.dumps(result, indent=2, default=str)

    @staticmethod
    def excel_to_csv_per_sheet(
        file_path: Union[str, Path],
        output_dir: Optional[Union[str, Path]] = None,
        include_sheet_names: Optional[List[str]] = None
    ) -> List[Path]:
        """
        Convert Excel sheets to separate CSV files.

        Args:
            file_path: Path to Excel file
            output_dir: Directory to save CSV files (None = same as Excel file)
            include_sheet_names: Sheets to include (None = all)

        Returns:
            List of paths to created CSV files

        Example:
            >>> csv_files = ExcelHandler.excel_to_csv_per_sheet("data.xlsx")
            >>> for csv in csv_files:
            ...     print(f"Created: {csv}")
        """
        file_path = Path(file_path)
        output_dir = Path(output_dir) if output_dir else file_path.parent
        output_dir.mkdir(parents=True, exist_ok=True)

        logger.info(f"Converting {file_path.name} sheets to CSV files")

        excel_file = pd.ExcelFile(file_path)
        sheet_names = include_sheet_names or excel_file.sheet_names

        csv_files = []

        for sheet_name in sheet_names:
            if sheet_name not in excel_file.sheet_names:
                continue

            df = pd.read_excel(excel_file, sheet_name=sheet_name)

            # Create safe filename
            safe_sheet_name = "".join(c if c.isalnum() else "_" for c in sheet_name)
            csv_filename = f"{file_path.stem}_{safe_sheet_name}.csv"
            csv_path = output_dir / csv_filename

            df.to_csv(csv_path, index=False)
            csv_files.append(csv_path)
            logger.info(f"Created CSV: {csv_path}")

        return csv_files

    @staticmethod
    def get_excel_summary(file_path: Union[str, Path]) -> Dict:
        """
        Get summary information about an Excel file.

        Args:
            file_path: Path to Excel file

        Returns:
            Dictionary with file summary information

        Example:
            >>> summary = ExcelHandler.get_excel_summary("data.xlsx")
            >>> print(f"Sheets: {summary['sheet_count']}")
        """
        file_path = Path(file_path)

        try:
            workbook = load_workbook(file_path, read_only=True)
            excel_file = pd.ExcelFile(file_path)

            sheets_info = []
            for sheet_name in workbook.sheetnames:
                df = pd.read_excel(excel_file, sheet_name=sheet_name)
                sheets_info.append({
                    "name": sheet_name,
                    "rows": len(df),
                    "columns": len(df.columns),
                    "column_names": list(df.columns)
                })

            return {
                "filename": file_path.name,
                "sheet_count": len(workbook.sheetnames),
                "sheet_names": workbook.sheetnames,
                "sheets": sheets_info,
                "file_size_mb": file_path.stat().st_size / (1024 * 1024)
            }

        except Exception as e:
            logger.error(f"Failed to get Excel summary: {e}")
            raise

    @staticmethod
    def prepare_excel_for_rag(
        file_path: Union[str, Path],
        output_dir: Optional[Union[str, Path]] = None,
        format: str = "markdown"
    ) -> Path:
        """
        Prepare Excel file for RAG ingestion by converting to optimal format.

        Args:
            file_path: Path to Excel file
            output_dir: Output directory (None = same as source)
            format: Output format ('markdown', 'text', 'json')

        Returns:
            Path to converted file

        Example:
            >>> converted = ExcelHandler.prepare_excel_for_rag(
            ...     "data.xlsx",
            ...     format="markdown"
            ... )
            >>> # Now upload the converted file to RAG system
        """
        file_path = Path(file_path)
        output_dir = Path(output_dir) if output_dir else file_path.parent
        output_dir.mkdir(parents=True, exist_ok=True)

        if format == "markdown":
            content = ExcelHandler.excel_to_markdown(file_path)
            output_path = output_dir / f"{file_path.stem}_converted.md"
            output_path.write_text(content, encoding="utf-8")

        elif format == "text":
            content = ExcelHandler.excel_to_structured_text(file_path)
            output_path = output_dir / f"{file_path.stem}_converted.txt"
            output_path.write_text(content, encoding="utf-8")

        elif format == "json":
            content = ExcelHandler.excel_to_json(file_path)
            output_path = output_dir / f"{file_path.stem}_converted.json"
            output_path.write_text(content, encoding="utf-8")

        else:
            raise ValueError(f"Unsupported format: {format}")

        logger.info(f"Prepared Excel for RAG: {output_path}")
        return output_path
